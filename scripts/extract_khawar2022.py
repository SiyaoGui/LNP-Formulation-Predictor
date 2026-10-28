"""从 Khawar 2022 的补充材料抽取脂质/水分配系数，产出 data/processed/partition_coefficients.csv。

来源: Khawar, M.I.; Mahmood, A.; Nabi, D. Scientific Reports 12(1), 14936, 2022.
      DOI 10.1038/s41598-022-19452-6，CC BY 4.0
授权登记: docs/data-licensing.md 第 2.1 节
表 schema: docs/schema-partition-coefficients.md

四类系数，用 coefficient_type 区分，不合并:
  logKlw    储存脂(三酰甘油)/水   Table S1 + SMILES 来自 S10/S11   n=305
  logKpw    磷脂/水               Table S2 + SMILES 来自 S12/S13   n=131
  logKlipw  脂质体/水             Table S15                        n=36
  logKtw    组织/水               Table S14                        n=18

Abraham 溶质描述符 E,S,A,B,V,L 只在 S1/S2 里有，一并抽出。

用法:
    python3 -I scripts/extract_khawar2022.py data/raw/khawar2022_klipw/41598_2022_19452_MOESM1_ESM.docx

原始文件是外部下载的数据，本脚本只读取，不改写。
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from rdkit import Chem, RDLogger
from rdkit.Chem import Descriptors

RDLogger.DisableLog("rdApp.*")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "processed" / "partition_coefficients.csv"
BASELINE = ROOT / "merged_formulations_v2.csv"

SOURCE_DOI = "10.1038/s41598-022-19452-6"
SOURCE_LICENSE = "CC-BY-4.0"
SOURCE_ATTRIBUTION = (
    "Khawar, Muhammad Irfan; Mahmood, Azhar; Nabi, Deedar. "
    "Exploring the role of octanol-water partition coefficient and Henry's law constant "
    "in predicting the lipid-water partition coefficients of organic chemicals. "
    "Scientific Reports 12(1), 14936, 2022. "
    "DOI: 10.1038/s41598-022-19452-6. Licensed under CC BY 4.0."
)

# 表索引 -> (原文表号, 说明)。索引由 docx 里表格出现顺序决定，已人工核对
T_S1_KLW = 0      # 305 行 logKlw 实测 + Abraham 描述符，无 SMILES
T_S2_KPW = 1      # 131 行 logKpw 实测 + Abraham 描述符，无 SMILES
T_S10_KLW_TRAIN = 9   # 245 行，带 SMILES
T_S11_KLW_VALID = 10  # 60 行，带 SMILES
T_S12_KPW_TRAIN = 11  # 107 行，带 SMILES
T_S13_KPW_VALID = 12  # 24 行，带 SMILES
T_S14_KTW = 13    # 18 行 logKtw 测试集，带 SMILES
T_S15_KLIPW = 14  # 36 行 logKlipw 测试集，带 SMILES


def table_to_df(doc, idx):
    """把 docx 表格转成 DataFrame，第一行当表头。"""
    tb = doc.tables[idx]
    rows = [[c.text.strip() for c in r.cells] for r in tb.rows]
    if not rows:
        return pd.DataFrame()
    header, body = rows[0], rows[1:]
    # 表头去重（docx 里合并单元格会产生重复列名）
    seen, cols = {}, []
    for h in header:
        if h in seen:
            seen[h] += 1
            cols.append(f"{h}__{seen[h]}")
        else:
            seen[h] = 0
            cols.append(h)
    return pd.DataFrame(body, columns=cols)


def norm_cas(x):
    """规范化 CAS：去空白，只留数字和连字符。空值返回 None。"""
    if not isinstance(x, str):
        return None
    s = re.sub(r"[^0-9\-]", "", x.strip())
    return s if s else None


def to_float(x):
    if not isinstance(x, str):
        return None
    s = x.strip().replace("−", "-").replace("–", "-")  # Unicode 减号
    try:
        return float(s)
    except ValueError:
        return None


def mol_info(smiles):
    """返回 (inchikey, mw, canonical_smiles)；无法解析返回三个 None。"""
    if not isinstance(smiles, str) or not smiles.strip():
        return None, None, None
    mol = Chem.MolFromSmiles(smiles.strip())
    if mol is None:
        return None, None, None
    try:
        key = Chem.MolToInchiKey(mol)
    except Exception:
        key = None
    return key, Descriptors.MolWt(mol), Chem.MolToSmiles(mol)


def build_smiles_lookup(doc, idxs):
    """从带 SMILES 的表里建 CAS -> SMILES 的查找表。"""
    lut, collisions = {}, 0
    for i in idxs:
        df = table_to_df(doc, i)
        for _, r in df.iterrows():
            cas = norm_cas(r.get("CAS No") or r.get("CAS-RN"))
            smi = (r.get("SMILES") or "").strip()
            if not cas or not smi:
                continue
            if cas in lut and lut[cas] != smi:
                collisions += 1
            lut[cas] = smi
    return lut, collisions


def extract_with_descriptors(doc, table_idx, coef_col, coef_type, smiles_lut, source_location):
    """抽 S1/S2：有 Abraham 描述符，SMILES 从查找表补。"""
    df = table_to_df(doc, table_idx)
    out = []
    for _, r in df.iterrows():
        cas = norm_cas(r.get("CAS No") or r.get("CAS-RN"))
        val = to_float(r.get(coef_col))
        if val is None:
            continue
        smi = smiles_lut.get(cas) if cas else None
        key, mw, canon = mol_info(smi)
        out.append(
            {
                "compound_name": (r.get("Chemicals") or "").strip(),
                "cas": cas,
                "smiles": canon,
                "smiles_as_reported": smi,
                "inchikey": key,
                "mw": mw,
                "coefficient_type": coef_type,
                "log_value": val,
                "log_kow": to_float(r.get("logKow")),
                "log_kaw": to_float(r.get("logKaw")),
                "abraham_E": to_float(r.get("E")),
                "abraham_S": to_float(r.get("S")),
                "abraham_A": to_float(r.get("A")),
                "abraham_B": to_float(r.get("B")),
                "abraham_V": to_float(r.get("V")),
                "abraham_L": to_float(r.get("L")),
                "id_source": "paper_smiles" if smi else "uncertain",
                "source_location": source_location,
            }
        )
    return out


def extract_test_table(doc, table_idx, coef_col, coef_type, source_location):
    """抽 S14/S15：自带 SMILES，无 Abraham 描述符。"""
    df = table_to_df(doc, table_idx)
    name_col = "Chemicals" if "Chemicals" in df.columns else "Compounds"
    out = []
    for _, r in df.iterrows():
        val = to_float(r.get(coef_col))
        if val is None:
            continue
        smi = (r.get("SMILES") or "").strip()
        key, mw, canon = mol_info(smi)
        out.append(
            {
                "compound_name": (r.get(name_col) or "").strip(),
                "cas": norm_cas(r.get("CAS No") or r.get("CAS-RN")),
                "smiles": canon,
                "smiles_as_reported": smi,
                "inchikey": key,
                "mw": mw,
                "coefficient_type": coef_type,
                "log_value": val,
                "log_kow": None,
                "log_kaw": None,
                "abraham_E": None, "abraham_S": None, "abraham_A": None,
                "abraham_B": None, "abraham_V": None, "abraham_L": None,
                "id_source": "paper_smiles" if smi else "uncertain",
                "source_location": source_location,
            }
        )
    return out


def baseline_inchikeys():
    """现有小分子子集的 InChIKey 集合。"""
    df = pd.read_csv(BASELINE, low_memory=False)
    large = {"mRNA", "siRNA", "DNA", "ASO", "protein", "mRNA, siRNA"}
    nondrug = {"empty", "gold_nanoparticles"}
    sm = df[
        df["payload_smiles"].notna()
        & (df["payload_smiles"].astype(str).str.strip() != "")
        & ~df["payload_type"].isin(large)
        & ~df["payload_type"].isin(nondrug)
    ]
    keys = set()
    for s in sm["payload_smiles"]:
        k, _, _ = mol_info(s)
        if k:
            keys.add(k)
    return keys


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx", help="Khawar 2022 补充材料 docx 的路径")
    args = ap.parse_args()

    path = Path(args.docx)
    if not path.is_file():
        print(f"找不到文件 {path}")
        return 1

    doc = Document(str(path))
    print(f"读入 {path.name}，表格 {len(doc.tables)} 个")

    lut_klw, c1 = build_smiles_lookup(doc, [T_S10_KLW_TRAIN, T_S11_KLW_VALID])
    lut_kpw, c2 = build_smiles_lookup(doc, [T_S12_KPW_TRAIN, T_S13_KPW_VALID])
    print(f"SMILES 查找表: logKlw {len(lut_klw)} 条, logKpw {len(lut_kpw)} 条")
    if c1 or c2:
        print(f"  注意: 同一 CAS 对应不同 SMILES 的冲突 {c1 + c2} 处，后出现的覆盖先出现的")

    records = []
    records += extract_with_descriptors(doc, T_S1_KLW, "logKlw", "logKlw", lut_klw, "Table S1")
    records += extract_with_descriptors(doc, T_S2_KPW, "logKpw", "logKpw", lut_kpw, "Table S2")
    records += extract_test_table(doc, T_S15_KLIPW, "logKlipw", "logKlipw", "Table S15")
    records += extract_test_table(doc, T_S14_KTW, "logKtw", "logKtw", "Table S14")

    df = pd.DataFrame(records)
    df.insert(0, "row_id", range(1, len(df) + 1))

    # schema 要求的固定列
    df["species"] = "unspecified"      # 原文未逐行给出带电态与 pH
    df["value_origin"] = "experimental"  # S1/S2/S14/S15 的表题均注明 experimental
    df["temperature_c"] = None
    df["ph"] = None
    df["ionic_strength"] = None
    df["measurement_method"] = None
    df["pubchem_cid"] = None
    df["source_doi"] = SOURCE_DOI
    df["source_license"] = SOURCE_LICENSE
    df["source_attribution"] = SOURCE_ATTRIBUTION
    df["primary_source_doi"] = None
    df["primary_source_available"] = False  # S1/S2 无逐行出处列
    df["qc_flag"] = "ok"
    df["qc_reason"] = None

    # 无法解析 SMILES 的标为 suspect，不静默删除
    bad = df["inchikey"].isna()
    df.loc[bad, "qc_flag"] = "suspect"
    df.loc[bad, "qc_reason"] = "SMILES 缺失或 RDKit 无法解析，无法计算 InChIKey"

    base = baseline_inchikeys()
    df["overlaps_repo_compound"] = df["inchikey"].isin(base) & df["inchikey"].notna()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False)

    print()
    print("=" * 68)
    print("抽取结果")
    print("=" * 68)
    for t, g in df.groupby("coefficient_type"):
        ok = g[g["qc_flag"] == "ok"]
        print(f"  {t:<10} 记录 {len(g):>4}  可解析 {len(ok):>4}  唯一化合物 {ok['inchikey'].nunique():>4}")
    print(f"  {'合计':<10} 记录 {len(df):>4}  可解析 {(~bad).sum():>4}  唯一化合物 {df.loc[~bad,'inchikey'].nunique():>4}")
    print()
    print(f"  SMILES 无法解析 {int(bad.sum())} 条，已标 qc_flag=suspect")
    print(f"  与现有 22 个化合物重叠 {int(df['overlaps_repo_compound'].sum())} 条记录，"
          f"{df.loc[df['overlaps_repo_compound'], 'inchikey'].nunique()} 个化合物")
    print()
    print(f"已写入 {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
